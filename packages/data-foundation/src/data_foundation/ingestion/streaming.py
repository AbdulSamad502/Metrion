"""Streaming CSV/XLSX readers (cfo-data-intelligence.md §6.3, D2, D5).

The old `read_csv_file` decoded the whole upload and held every row as a dict (~31x the file size). These
readers yield bounded `RawFile` chunks instead, so memory is set by `chunk_rows`, not by file size, and
they emit exactly the `RawFile` shape the existing adapters already consume (dict rows, or ordered
(header, value) pairs when a header appears twice — the existing duplicate-header convention).

Safety (spec §8): XLSX is opened read-only, values only (formulas are never evaluated, macros and external
links are ignored), after a zip-structure check that rejects zip bombs and oversized archives.
"""

from __future__ import annotations

import codecs
import csv
import zipfile
from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from fractions import Fraction
from pathlib import Path
from typing import Any

from data_foundation.types import RawFile, RowData


class StreamError(Exception):
    """A file that cannot be streamed. `code` is stable for the caller/UI (spec E11/E12/AC-D10)."""

    def __init__(self, code: str, detail: str) -> None:
        self.code = code
        self.detail = detail
        super().__init__(f"{code}: {detail}")


@dataclass(frozen=True, slots=True)
class StreamLimits:
    """Configuration — supplied by the caller (CLAUDE.md: thresholds are config, not constants)."""

    chunk_rows: int
    max_rows: int
    header_scan_rows: int
    min_header_fill: Fraction  # share of a row's cells that must be non-empty text for it to be the header row
    max_uncompressed_bytes: int  # XLSX: total uncompressed size of the archive
    max_zip_ratio: int  # XLSX: uncompressed/compressed ratio of any single member
    max_zip_members: int
    max_sheets: int


@dataclass(frozen=True, slots=True)
class StreamHeader:
    columns: tuple[str, ...]
    header_row_index: int  # 0-based row index of the header in the source
    sheet_name: str | None
    has_duplicate_headers: bool
    encoding: str | None  # CSV only


@dataclass(slots=True)
class StreamStats:
    """Filled as the stream is consumed — what the data-quality report needs."""

    rows_read: int = 0
    blank_rows_skipped: int = 0
    ragged_rows: int = 0  # rows whose cell count differed from the header (padded/truncated, never dropped)


def _cell_to_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return format(Decimal(repr(v)), "f")  # no scientific notation, no float rounding artefacts
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0, 0) else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _is_header_like(cells: Sequence[str], width: int, min_fill: Fraction) -> bool:
    if width == 0:
        return False
    nonempty = [c for c in cells if c.strip()]
    if Fraction(len(nonempty), width) < min_fill:
        return False
    # a header row is (almost) all text; a row with numbers in it is data or a summary line
    texty = sum(1 for c in nonempty if not _looks_numeric(c))
    return texty * 5 >= len(nonempty) * 4


def _looks_numeric(s: str) -> bool:
    t = s.replace(",", "").replace("₹", "").strip()
    try:
        Decimal(t)
        return True
    except Exception:  # noqa: BLE001 - any parse failure just means "not a number"
        return False


def _find_header(rows: Sequence[Sequence[str]], limits: StreamLimits) -> int | None:
    width = max((len(r) for r in rows), default=0)
    for i, r in enumerate(rows):
        if _is_header_like(r, width, limits.min_header_fill):
            return i
    return None


class RowStream:
    """A resolved header plus a one-shot generator of bounded chunks. Consume `chunks()` once."""

    def __init__(self, header: StreamHeader, rows: Iterator[Sequence[str]], limits: StreamLimits) -> None:
        self.header = header
        self.stats = StreamStats()
        self._rows = rows
        self._limits = limits
        self._width = len(header.columns)

    def _shape(self, cells: Sequence[str]) -> RowData:
        if len(cells) != self._width:
            self.stats.ragged_rows += 1
        padded = list(cells[: self._width]) + [""] * max(0, self._width - len(cells))
        if self.header.has_duplicate_headers:
            return tuple(zip(self.header.columns, padded))
        return dict(zip(self.header.columns, padded))

    def close(self) -> None:
        """Release the underlying file/workbook if the stream was opened but not (fully) consumed."""
        closer = getattr(self._rows, "close", None)
        if closer is not None:
            closer()

    def chunks(self) -> Iterator[RawFile]:
        buf: list[RowData] = []
        for cells in self._rows:
            if not any(c.strip() for c in cells):
                self.stats.blank_rows_skipped += 1
                continue
            self.stats.rows_read += 1
            if self.stats.rows_read > self._limits.max_rows:
                raise StreamError("too_large", f"the file has more than the allowed {self._limits.max_rows} data rows")
            buf.append(self._shape(cells))
            if len(buf) >= self._limits.chunk_rows:
                yield RawFile(columns=self.header.columns, rows=buf)
                buf = []
        if buf:
            yield RawFile(columns=self.header.columns, rows=buf)


# ------------------------------------------------------------------------------------------------ CSV
def _detect_encoding(path: Path) -> str:
    with path.open("rb") as fh:
        head = fh.read(1024 * 1024)
    try:
        codecs.getdecoder("utf-8-sig")(head, "strict")
        return "utf-8-sig"
    except UnicodeDecodeError as exc:
        # A multi-byte sequence cut at the 1 MiB boundary is not a real decode error.
        if exc.start >= len(head) - 4 and len(head) == 1024 * 1024:
            return "utf-8-sig"
        return "cp1252"


def open_csv(path: Path, limits: StreamLimits) -> RowStream:
    if path.stat().st_size == 0:
        raise StreamError("empty_file", "the file is empty")
    enc = _detect_encoding(path)
    fh = path.open("r", encoding=enc, newline="", errors="replace")
    reader = csv.reader(fh)
    head: list[list[str]] = []
    for row in reader:
        head.append([c.strip() for c in row])
        if len(head) >= limits.header_scan_rows:
            break
    idx = _find_header(head, limits)
    if idx is None:
        fh.close()
        raise StreamError("no_header", "no header row could be identified in the first rows of the file")
    columns = tuple(head[idx])
    header = StreamHeader(columns=columns, header_row_index=idx, sheet_name=None,
                          has_duplicate_headers=len(set(columns)) != len(columns), encoding=enc)

    def _rows() -> Iterator[Sequence[str]]:
        try:
            yield from head[idx + 1:]
            for row in reader:
                yield [c.strip() for c in row]
        finally:
            fh.close()

    return RowStream(header, _rows(), limits)


# ------------------------------------------------------------------------------------------------ XLSX
def check_xlsx_archive(path: Path, limits: StreamLimits) -> None:
    """Reject zip bombs and oversized archives *before* any parser opens them (AC-D10)."""
    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile:
        raise StreamError("not_xlsx", "the file is not a valid .xlsx workbook") from None
    with zf:
        infos = zf.infolist()
        if len(infos) > limits.max_zip_members:
            raise StreamError("unsafe_archive", f"the archive has {len(infos)} parts (limit {limits.max_zip_members})")
        total = 0
        for info in infos:
            total += info.file_size
            if info.compress_size > 0 and info.file_size // info.compress_size > limits.max_zip_ratio:
                raise StreamError("unsafe_archive", f"part {info.filename!r} expands more than {limits.max_zip_ratio}x (possible zip bomb)")
            if info.compress_size == 0 and info.file_size > 0:
                raise StreamError("unsafe_archive", f"part {info.filename!r} has an impossible size")
        if total > limits.max_uncompressed_bytes:
            raise StreamError("unsafe_archive", f"the archive would expand to {total} bytes (limit {limits.max_uncompressed_bytes})")
        if "xl/workbook.xml" not in {i.filename for i in infos}:
            raise StreamError("not_xlsx", "the archive is not an Excel workbook")


def _open_workbook(path: Path, limits: StreamLimits) -> Any:
    check_xlsx_archive(path, limits)
    import openpyxl  # type: ignore[import-untyped]

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    if len(wb.sheetnames) > limits.max_sheets:
        wb.close()
        raise StreamError("unsafe_archive", f"the workbook has {len(wb.sheetnames)} sheets (limit {limits.max_sheets})")
    return wb


def _sheet_rows(ws: Any) -> Iterator[list[str]]:
    # Some exporters (the Meesho GST workbooks among them) write a sheet dimension of A1 although the sheet holds many columns; openpyxl's read-only mode then
    # returns only column A. Forget the declared dimension and let the reader find the real extent.
    reset = getattr(ws, "reset_dimensions", None)
    if callable(reset):
        reset()
    for row in ws.iter_rows(values_only=True):
        yield [_cell_to_str(v) for v in row]


def list_xlsx_sheets(path: Path, limits: StreamLimits) -> list[tuple[str, bool]]:
    """(sheet name, has a detectable header row) for each visible sheet — so the seller can choose when
    several candidates exist (spec E3)."""
    wb = _open_workbook(path, limits)
    try:
        out: list[tuple[str, bool]] = []
        for ws in wb.worksheets:
            if getattr(ws, "sheet_state", "visible") != "visible":
                continue
            head: list[list[str]] = []
            for r in _sheet_rows(ws):
                head.append(r)
                if len(head) >= limits.header_scan_rows:
                    break
            out.append((ws.title, _find_header(head, limits) is not None))
        return out
    finally:
        wb.close()


def open_xlsx(path: Path, limits: StreamLimits, sheet_name: str | None = None) -> RowStream:
    wb = _open_workbook(path, limits)
    chosen = None
    for ws in wb.worksheets:
        if getattr(ws, "sheet_state", "visible") != "visible":
            continue
        if sheet_name is not None and ws.title != sheet_name:
            continue
        chosen = ws
        break
    if chosen is None:
        wb.close()
        raise StreamError("no_sheet", f"sheet {sheet_name!r} was not found" if sheet_name else "the workbook has no visible sheet")

    it = _sheet_rows(chosen)
    head: list[list[str]] = []
    for r in it:
        head.append(r)
        if len(head) >= limits.header_scan_rows:
            break
    idx = _find_header(head, limits)
    if idx is None:
        wb.close()
        raise StreamError("no_header", f"no header row could be identified on sheet {chosen.title!r}")
    columns = tuple(c.strip() for c in head[idx])
    header = StreamHeader(columns=columns, header_row_index=idx, sheet_name=chosen.title,
                          has_duplicate_headers=len(set(columns)) != len(columns), encoding=None)

    def _rows() -> Iterator[Sequence[str]]:
        try:
            yield from head[idx + 1:]
            yield from it
        finally:
            wb.close()

    return RowStream(header, _rows(), limits)


def open_report(path: Path, limits: StreamLimits, *, sheet_name: str | None = None) -> RowStream:
    """Open by extension. `.xlsm`/`.xlsb`/`.xls` are not accepted (macros and legacy binary formats)."""
    suffix = path.suffix.lower()
    if suffix in (".csv", ".txt"):
        return open_csv(path, limits)
    if suffix == ".xlsx":
        return open_xlsx(path, limits, sheet_name)
    raise StreamError("unsupported_type", f"{suffix or 'this file type'} is not supported; upload .csv or .xlsx")

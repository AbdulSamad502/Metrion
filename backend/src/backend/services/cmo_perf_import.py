"""CMO 3.0 M8: importing a seller's own marketplace performance report (CSV or XLSX). Zero AI credits; no model; nothing is sent anywhere.

Two steps, both pure reads of the same bytes: `preview` shows exactly what would be saved (and every reason it cannot be), `commit` saves it only if the file's SHA-256 is the one the seller previewed.
Rows are only ever added: a record that already exists (same marketplace, SKU and date) is never overwritten, and the same file cannot be imported twice. A blank cell stays blank (NULL)."""

from __future__ import annotations

import csv
import hashlib
import io
from datetime import date
from typing import Any
from uuid import UUID

from sqlalchemy.exc import IntegrityError

from backend.repositories import cmo_performance_repo as repo
from backend.services.cmo_performance import PerformanceRefused, _mp
from slim_cmo.history import metrics as mx
from slim_cmo.history import report_import as ri

MAX_BYTES = 10 * 1024 * 1024
MAX_SHEET_ROWS = ri.MAX_ROWS + 5
_ZIP = b"PK\x03\x04"


def read_table(name: str, data: bytes) -> list[list[Any]]:
    """The file's first sheet / the CSV as rows of cell values. Cell values only: a workbook's formulas are never evaluated."""
    low = (name or "").casefold()
    if not data:
        raise PerformanceRefused("EMPTY_FILE", "The file is empty.")
    if len(data) > MAX_BYTES:
        raise PerformanceRefused("TOO_LARGE", f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB. Split it by period.", 413)
    is_zip = data[:4] == _ZIP
    if low.endswith(".xlsx") or (is_zip and not low.endswith(".csv")):
        if not is_zip:
            raise PerformanceRefused("BAD_FILE", "This is not a valid .xlsx workbook.")
        try:
            from openpyxl import load_workbook  # type: ignore[import-untyped]  # noqa: PLC0415

            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            out: list[list[Any]] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_SHEET_ROWS:
                    break
                out.append(list(row))
            return out
        except Exception as exc:  # noqa: BLE001
            raise PerformanceRefused("BAD_FILE", "This workbook could not be read. Save it again as .xlsx or .csv and retry.") from exc
    if not (low.endswith(".csv") or low.endswith(".txt") or "." not in low):
        raise PerformanceRefused("BAD_FILE", "Only .csv and .xlsx files can be imported.")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = data.decode("cp1252")
        except UnicodeDecodeError:
            raise PerformanceRefused("BAD_FILE", "This file is not readable text. Save it as UTF-8 CSV or .xlsx.") from None
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)][:MAX_SHEET_ROWS]


def _display(m: str, v: int | None) -> str:
    return "—" if v is None else mx.fmt_money(v) if m == "sales" else f"{v:,}"


def preview(tenant: UUID, marketplace: str, name: str, data: bytes, today: date) -> dict[str, Any]:
    """What importing this file would do. Saves nothing."""
    return _preview(tenant, marketplace, name, data, today)[0]


def _preview(tenant: UUID, marketplace: str, name: str, data: bytes, today: date) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    mp = _mp(marketplace)
    sha = hashlib.sha256(data).hexdigest()
    table = read_table(name, data)
    v = ri.validate(table, mp, today)
    earlier = repo.import_by_sha(tenant, sha)
    out: dict[str, Any] = {"sha256": sha, "fileName": name, "marketplace": mp, "ok": v["ok"], "errors": v["errors"], "errorCount": v["errorCount"], "warnings": v["warnings"][:30],
                           "metrics": v["metrics"], "ignoredColumns": v["ignoredColumns"], "rowCount": v["rowCount"], "period": None, "newRows": 0, "duplicateRows": 0, "conflicts": [], "sample": [],
                           "alreadyImported": None, "canImport": False, "aiCredits": 0}
    if earlier:
        out["alreadyImported"] = {"importedAt": earlier["importedAt"].isoformat(), "fileName": earlier["sourceName"], "periodStart": earlier["periodStart"].isoformat(), "periodEnd": earlier["periodEnd"].isoformat()}
    if not v["ok"]:
        return out, []
    lo, hi = v["period"]
    out["period"] = {"start": lo.isoformat(), "end": hi.isoformat(), "days": (hi - lo).days + 1}
    have = repo.existing_lines(tenant, mp, sorted({x["sku"] for x in v["lines"]}), lo, hi)
    new, dup, conflicts = [], 0, []
    for x in v["lines"]:
        was = have.get((x["sku"], x["date"]))
        if was is None:
            new.append(x)
        elif all(was.get(m) == x.get(m) for m in v["metrics"]):
            dup += 1
        else:
            conflicts.append({"row": x["row"], "sku": x["sku"], "date": x["date"].isoformat(),
                              "message": f"Row {x['row']}: {x['sku']} on {x['date'].isoformat()} is already saved with different numbers. Existing records are never overwritten."})
    out.update(newRows=len(new), duplicateRows=dup, conflicts=conflicts[:ri.MAX_SHOWN], conflictCount=len(conflicts))
    out["sample"] = [{"date": x["date"].isoformat(), "sku": x["sku"], **{m: _display(m, x.get(m)) for m in v["metrics"]}} for x in new[:10]]
    out["canImport"] = bool(new) and not conflicts and earlier is None
    if not new and not conflicts:
        out["warnings"].append("Every row in this file is already saved, so there is nothing to import.")
    return out, new


def commit(tenant: UUID, user: UUID, marketplace: str, name: str, data: bytes, expected_sha256: str, today: date) -> dict[str, Any]:
    p, new = _preview(tenant, marketplace, name, data, today)
    if p["sha256"] != expected_sha256:
        raise PerformanceRefused("FILE_CHANGED", "This is not the file you previewed. Preview it again before importing.", 409)
    if p["alreadyImported"]:
        raise PerformanceRefused("ALREADY_IMPORTED", f"This exact file was already imported on {p['alreadyImported']['importedAt'][:10]}. Nothing was changed.", 409)
    if not p["ok"]:
        raise PerformanceRefused("INVALID_FILE", f"The file has {p['errorCount']} problem(s). Nothing was saved. Fix them and preview again.")
    if p.get("conflictCount"):
        raise PerformanceRefused("CONFLICTING_RECORDS", "Some rows are already saved with different numbers. Existing records are never overwritten, so nothing was saved.", 409)
    if not p["newRows"]:
        raise PerformanceRefused("NOTHING_NEW", "Every row in this file is already saved. Nothing was imported.", 409)
    lo, hi = date.fromisoformat(p["period"]["start"]), date.fromisoformat(p["period"]["end"])
    skipped = int(p["rowCount"]) - int(p["newRows"])
    try:
        saved = repo.save_import(tenant, user, marketplace=p["marketplace"], source_name=name[:200] or "report", sha256=p["sha256"], period=(lo, hi), metrics=p["metrics"], lines=new, skipped=skipped)
    except IntegrityError:  # the same file was imported a moment ago in another request: nothing is added twice
        raise PerformanceRefused("ALREADY_IMPORTED", "This exact file was already imported. Nothing was changed.", 409) from None
    return {"imported": True, "importId": saved["id"], "importedAt": saved["importedAt"].isoformat(), "marketplace": p["marketplace"], "rowsSaved": p["newRows"], "rowsSkipped": skipped,
            "period": p["period"], "metrics": p["metrics"], "aiCredits": 0,
            "message": f"Imported {p['newRows']} row(s) for {lo.isoformat()} to {hi.isoformat()}. Nothing was published or changed on a marketplace."}

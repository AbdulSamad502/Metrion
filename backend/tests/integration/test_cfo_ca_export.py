"""The CA export pack, real Postgres (docs/specs/cfo-phase3-scope.md AC-E1…E6). The workbook is opened and read back; totals are compared with the planner's own results."""

from __future__ import annotations

import io
import uuid
import zipfile
from datetime import date, datetime, timezone

import pytest
from openpyxl import load_workbook

from calc_engine.expense_ledger import NewEntry
from calc_engine.formulas.expense_accounting import ExpenseClass
from calc_engine.repositories import expense_ledger_repo as ledger
from cfo.planner.plan import PeriodRef

from backend.services import cfo_ca_export
from tests.integration.test_cfo_expense_ledger import _register_owner, app_client, fake_clerk, _seed_zero_cap  # noqa: F401
from tests.integration.test_cfo_query_planner import SEP, _cogs, _file_and_run, _line, plan
from tests.integration.test_cfo_ads import ad_lines, ask, by_metric

USER = uuid.uuid4()


def book() -> uuid.UUID:
    t = uuid.uuid4()
    run = _file_and_run(t)
    for d in (3, 4, 5, 6):
        _line(t, run, day=date(2026, 9, d), sku="A", gross=50_000, fee=5_000)
    _cogs(t, "A", 4_000)
    return t


def cell_map(pack: cfo_ca_export.CaPack, sheet: str) -> dict[str, list]:
    sh = next(s for s in pack.sheets if s.name == sheet)
    return {str(r[0]): r for r in sh.rows}


def test_the_pack_has_the_fixed_sheets_with_the_disclaimer_first_and_engine_numbers() -> None:
    t = book()
    ledger.add_entries(t, [NewEntry(entry_date=date(2026, 9, 10), amount_paise=30_000, category="Rent", vendor="L")], USER)
    ledger.confirm_class(t, klass=ExpenseClass.OVERHEAD, confirmed_by=USER, category="Rent")
    pack = cfo_ca_export.build_pack(t, USER, SEP)
    assert tuple(s.name for s in pack.sheets) == cfo_ca_export.SHEETS
    pnl = cell_map(pack, "P&L")
    cp_planner = by_metric(ask(t, [{"op": "metric", "params": {"metric": "CONTRIBUTION_PROFIT"}}]))["CONTRIBUTION_PROFIT"]
    cp_row = pnl["Contribution Profit (not accounting net profit)"]
    assert cp_row[1] == cp_planner.value_paise and cp_row[4] and cp_row[5]  # same number, with a ref and a formula version
    nrr_row = pnl["Net Realized Revenue"]
    assert nrr_row[1] == 200_000 and nrr_row[2] == "₹2,000.00"
    op_row = pnl["Management operating profit (before depreciation, interest and tax)"]
    assert op_row[1] == cp_planner.value_paise - 30_000 and op_row[3] == "DERIVED"
    exp = cell_map(pack, "Expenses")
    assert exp["Rent"][1] == "OVERHEAD" and exp["Rent"][2] == 30_000 and exp["Rent"][3] == "₹300.00"


def test_components_add_up_to_contribution_profit_in_the_pack() -> None:
    t = book()
    pnl = cell_map(cfo_ca_export.build_pack(t, USER, SEP), "P&L")
    nrr = pnl["Net Realized Revenue"][1]
    costs = sum(v[1] for k, v in pnl.items() if k.endswith("(cost, shown positive)"))
    assert nrr - costs == pnl["Contribution Profit (not accounting net profit)"][1]


def test_unavailable_items_are_listed_with_reasons_never_omitted_or_zero() -> None:
    t = uuid.uuid4()
    run = _file_and_run(t)
    _line(t, run, day=date(2026, 9, 5), sku="NOCOST", gross=50_000, fee=5_000)  # no cost price: contribution profit cannot be computed
    pack = cfo_ca_export.build_pack(t, USER, SEP)
    pnl = cell_map(pack, "P&L")
    cp = pnl["Contribution Profit (not accounting net profit)"]
    assert cp[1] is None and cp[2] == "UNAVAILABLE" and cp[3] == "UNAVAILABLE"
    op = pnl["Management operating profit (before depreciation, interest and tax)"]
    assert op[1] is None and op[3] == "UNAVAILABLE"
    ex = cell_map(pack, "Exclusions & unavailable")
    assert any(r[1] == "UNAVAILABLE" for r in ex.values())
    assert cell_map(pack, "Open data requests")  # has rows (the missing cost price request, or the "none open" marker)


def test_held_and_unclassified_expenses_are_visible_and_excluded() -> None:
    t = book()
    ledger.add_entries(t, [NewEntry(entry_date=date(2026, 9, 10), amount_paise=30_000, category="Misc", vendor="L"),
                           NewEntry(entry_date=date(2026, 9, 11), amount_paise=7_000, category="Marketplace commission", vendor="M")], USER)
    pack = cfo_ca_export.build_pack(t, USER, SEP)
    exp = cell_map(pack, "Expenses")
    assert exp["Misc"][5] == "UNCLASSIFIED" and any(str(k).startswith("(held: may overlap unconfirmed)") for k in exp)
    prov = [r for r in next(s for s in pack.sheets if s.name == "Provenance").rows if r[0] == "EXPENSE ENTRY"]
    assert len(prov) == 2 and any("HELD" in str(r[2]) for r in prov)


def test_the_tcs_tds_schedule_lists_levies_totals_them_by_code_and_says_they_are_not_costs() -> None:
    from data_foundation.repositories.tenant_context import tenant_scope
    from sqlalchemy import text

    t = book()
    with tenant_scope(t) as conn:
        sid = conn.execute(text("SELECT id FROM sales_line WHERE tenant_id = :t ORDER BY order_ref LIMIT 2"), {"t": str(t)}).all()
        for (lid,), amount in zip(sid, (1_200, 800)):
            conn.execute(text("INSERT INTO line_component (id, tenant_id, sales_line_id, settlement_line_id, component_kind, component_code, source_side, amount, source_column, is_verified) "
                              "VALUES (:i, :t, :l, NULL, 'LEVY', 'TCS', 'INVOICE', :a, 'x', false)"), {"i": str(uuid.uuid4()), "t": str(t), "l": str(lid), "a": amount})
    tax = cell_map(cfo_ca_export.build_pack(t, USER, SEP), "TCS-TDS schedule")
    assert tax["TOTAL TCS"][5] == 2_000 and tax["TOTAL TCS"][6] == "₹20.00"
    assert any("not costs" in k for k in tax)


def test_reconciliation_totals_exclude_amazon_and_name_it_withheld() -> None:
    from tests.integration.test_cfo_leak_claims import finding

    t = book()
    finding(t, marketplace="MEESHO")
    finding(t, marketplace="AMAZON")
    rec = next(s for s in cfo_ca_export.build_pack(t, USER, SEP).sheets if s.name == "Reconciliation")
    rows = rec.rows
    meesho = [r for r in rows if r[1] == "MEESHO" and not str(r[0]).startswith("TOTAL")]
    amazon = [r for r in rows if r[1] == "AMAZON" and not str(r[0]).startswith("TOTAL")]
    assert meesho[0][4] == 5_000 and meesho[0][6] == "NEW" and amazon[0][4] is None and amazon[0][5] == "WITHHELD" and amazon[0][6] == "WITHHELD"
    totals = [r for r in rows if str(r[0]).startswith("TOTAL")]
    assert [(r[1], r[4]) for r in totals] == [("MEESHO", 5_000)] and "not an amount recovered" in totals[0][7]


def test_advertising_report_spend_shows_in_the_pnl_components() -> None:
    t = book()
    run = _file_and_run(t)
    ad_lines(t, run)
    pnl = cell_map(cfo_ca_export.build_pack(t, USER, SEP), "P&L")
    assert pnl["Advertising costs (cost, shown positive)"][1] == 30_000


def test_xlsx_and_csv_zip_render_every_sheet_with_the_disclaimer_and_integer_paise() -> None:
    t = book()
    pack = cfo_ca_export.build_pack(t, USER, SEP)
    wb = load_workbook(io.BytesIO(cfo_ca_export.to_xlsx(pack)))
    assert wb.sheetnames == [s.name[:31] for s in pack.sheets]
    for ws in wb.worksheets:
        assert "NOT a tax return" in ws["A1"].value and ws["A3"].value
    pnl = wb["P&L"]
    money = [pnl.cell(row=r, column=2).value for r in range(4, pnl.max_row + 1) if pnl.cell(row=r, column=2).value is not None]
    assert money and all(isinstance(v, int) and not isinstance(v, bool) for v in money)  # integer paise, never a float
    z = zipfile.ZipFile(io.BytesIO(cfo_ca_export.to_csv_zip(pack)))
    assert len(z.namelist()) == len(pack.sheets) and all("NOT a tax return" in z.read(n).decode("utf-8").splitlines()[0] for n in z.namelist())


def test_the_export_is_tenant_isolated() -> None:
    a, b = book(), uuid.uuid4()
    ledger.add_entries(a, [NewEntry(entry_date=date(2026, 9, 10), amount_paise=30_000, category="Rent", vendor="L")], USER)
    pb = cfo_ca_export.build_pack(b, USER, SEP)
    assert cell_map(pb, "P&L")["Net Realized Revenue"][1] is None
    assert not any(r[0] == "EXPENSE ENTRY" for r in next(s for s in pb.sheets if s.name == "Provenance").rows)


def test_route_returns_a_workbook_or_a_zip_with_gates(app_client, fake_clerk, two_tenants, migrate_engine):
    tenant_a, _ = two_tenants
    h = _register_owner(fake_clerk, tenant_a)
    q = {"period_start": "2026-09-01", "period_end": "2026-09-30", "as_of_date": "2026-10-01"}
    r = app_client.get("/api/v1/cfo/ca-export", params=q, headers=h)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats") and "metrion-ca-pack-2026-09-01-to-2026-09-30.xlsx" in r.headers["content-disposition"]
    assert load_workbook(io.BytesIO(r.content)).sheetnames[0] == "Cover"
    z = app_client.get("/api/v1/cfo/ca-export", params={**q, "format": "zip"}, headers=h)
    assert z.status_code == 200 and z.headers["content-type"] == "application/zip" and zipfile.ZipFile(io.BytesIO(z.content)).namelist()
    assert app_client.get("/api/v1/cfo/ca-export", params={**q, "format": "pdf"}, headers=h).status_code == 400
    assert app_client.get("/api/v1/cfo/ca-export", params={**q, "period_end": "2026-08-01"}, headers=h).status_code == 400
    _seed_zero_cap(migrate_engine, agent_code="CFO", meter_code="cfo.ca_export")
    assert app_client.get("/api/v1/cfo/ca-export", params=q, headers=h).status_code == 402
